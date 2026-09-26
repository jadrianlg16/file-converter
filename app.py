"""Legacy desktop tool: Markdown -> DOCX with a small tkinter window.

Independent of the web service (web_app.py) and its converters package.
Needs: markdown, python-docx, beautifulsoup4 (see requirements.txt).
Run: python app.py
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import threading


def convert_md_to_docx(input_path, output_path):
    import markdown
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from bs4 import BeautifulSoup, NavigableString

    with open(input_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    html = markdown.markdown(
        md_content,
        extensions=["tables", "fenced_code", "nl2br"],
    )

    doc = Document()
    soup = BeautifulSoup(html, "html.parser")

    def add_runs(paragraph, element):
        for child in element.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text:
                    paragraph.add_run(text)
            elif child.name in ("strong", "b"):
                run = paragraph.add_run(child.get_text())
                run.bold = True
            elif child.name in ("em", "i"):
                run = paragraph.add_run(child.get_text())
                run.italic = True
            elif child.name == "code":
                run = paragraph.add_run(child.get_text())
                run.font.name = "Courier New"
                run.font.size = Pt(10)
            elif child.name == "a":
                run = paragraph.add_run(child.get_text())
                run.underline = True
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            else:
                add_runs(paragraph, child)

    for element in soup.children:
        if isinstance(element, NavigableString):
            continue
        tag = element.name

        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            p = doc.add_heading(level=int(tag[1]))
            p.clear()
            add_runs(p, element)

        elif tag == "p":
            p = doc.add_paragraph()
            add_runs(p, element)

        elif tag == "ul":
            for li in element.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, li)

        elif tag == "ol":
            for li in element.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Number")
                add_runs(p, li)

        elif tag == "pre":
            code_el = element.find("code")
            text = code_el.get_text() if code_el else element.get_text()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(10)

        elif tag == "blockquote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            add_runs(p, element)

        elif tag == "hr":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "auto")
            pBdr.append(bottom)
            pPr.append(pBdr)

        elif tag == "table":
            rows = element.find_all("tr")
            if rows:
                max_cols = max(len(r.find_all(["th", "td"])) for r in rows)
                table = doc.add_table(rows=len(rows), cols=max_cols)
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    cells = row.find_all(["th", "td"])
                    for j, cell in enumerate(cells):
                        if j < max_cols:
                            table.cell(i, j).text = cell.get_text()

    doc.save(output_path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("MD → DOCX Converter")
        self.resizable(False, False)
        self._build_ui()

    def _build_ui(self):
        frame = ttk.Frame(self, padding=16)
        frame.grid(sticky="nsew")

        ttk.Label(frame, text="Input (.md)").grid(row=0, column=0, sticky="w", pady=6)
        self.input_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.input_var, width=48).grid(
            row=0, column=1, padx=8, pady=6
        )
        ttk.Button(frame, text="Browse…", command=self._pick_input).grid(
            row=0, column=2, pady=6
        )

        ttk.Label(frame, text="Output (.docx)").grid(row=1, column=0, sticky="w", pady=6)
        self.output_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.output_var, width=48).grid(
            row=1, column=1, padx=8, pady=6
        )
        ttk.Button(frame, text="Browse…", command=self._pick_output).grid(
            row=1, column=2, pady=6
        )

        self.convert_btn = ttk.Button(
            frame, text="Convert", command=self._convert
        )
        self.convert_btn.grid(row=2, column=0, columnspan=3, pady=(14, 6))

        self.status_var = tk.StringVar(value="Ready.")
        ttk.Label(frame, textvariable=self.status_var, foreground="gray").grid(
            row=3, column=0, columnspan=3, pady=(0, 4)
        )

    def _pick_input(self):
        path = filedialog.askopenfilename(
            filetypes=[("Markdown files", "*.md *.markdown"), ("All files", "*.*")]
        )
        if path:
            self.input_var.set(path)
            if not self.output_var.get():
                self.output_var.set(os.path.splitext(path)[0] + ".docx")

    def _pick_output(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
        )
        if path:
            self.output_var.set(path)

    def _convert(self):
        inp = self.input_var.get().strip()
        out = self.output_var.get().strip()
        if not inp:
            messagebox.showerror("Error", "Please select an input file.")
            return
        if not out:
            messagebox.showerror("Error", "Please specify an output file.")
            return
        if not os.path.isfile(inp):
            messagebox.showerror("Error", f"Input file not found:\n{inp}")
            return

        self.convert_btn.config(state="disabled")
        self.status_var.set("Converting…")

        def run():
            try:
                convert_md_to_docx(inp, out)
                self.after(
                    0,
                    lambda: self.status_var.set(
                        f"Done! Saved to {os.path.basename(out)}"
                    ),
                )
                self.after(
                    0, lambda: messagebox.showinfo("Success", f"Saved to:\n{out}")
                )
            except Exception as e:
                # Bind the message now: Python unbinds `e` when this block
                # exits, before Tk gets around to running the callbacks.
                msg = str(e)
                self.after(0, lambda: self.status_var.set(f"Error: {msg}"))
                self.after(
                    0, lambda: messagebox.showerror("Conversion failed", msg)
                )
            finally:
                self.after(0, lambda: self.convert_btn.config(state="normal"))

        threading.Thread(target=run, daemon=True).start()


if __name__ == "__main__":
    App().mainloop()
