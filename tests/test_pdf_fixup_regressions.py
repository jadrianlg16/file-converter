"""Focused regressions for pdf_docx_fixup defects found in review.

Each test builds a tiny fixture with PyMuPDF / python-docx (the way
pdf_fixtures.py does) and pins one repaired behavior. None needs LibreOffice.
"""

import io
import os
import sys
from itertools import pairwise

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from converters import documents
from converters import pdf_docx_fixup as fx

fitz = pytest.importorskip("fitz")
docx = pytest.importorskip("docx")
pytest.importorskip("pdf2docx")

from docx.shared import Pt  # noqa: E402 - after the importorskip checks

import pdf_fixtures  # noqa: E402 - after the importorskip checks

NS = fx.W_NS
W, H = 595, 842


def _png(rgb=(0, 0, 0), size=(20, 20)) -> bytes:
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, *size), False)
    pix.set_rect(pix.irect, rgb)
    return pix.tobytes("png")


def _body_lines(page, pg):
    y = 150
    for i in range(10):
        page.insert_text((40, y), f"Parrafo {pg}.{i} del cuerpo del texto.", fontsize=10)
        y += 14


def _pdf2docx(pdf, out):
    from pdf2docx import Converter

    cv = Converter(pdf)
    try:
        cv.convert(out)
    finally:
        cv.close()
    return docx.Document(out)


def _drawings(el) -> int:
    return sum(1 for _ in el.iter(f"{NS}drawing"))


def _label_value_pdf(path):
    """Bold red label + regular red value; PyMuPDF keeps the word space
    inside the value span (' Juan Perez'), then a far-right date."""
    doc = fitz.open()
    page = doc.new_page()
    red = (0.8, 0, 0)
    page.insert_text((40, 100), "Nombre:", fontsize=10, fontname="hebo", color=red)
    x = 40 + fitz.get_text_length("Nombre: ", fontname="hebo", fontsize=10)
    page.insert_text((x, 100), "Juan Perez", fontsize=10, color=red)
    page.insert_text((400, 100), "Fecha: 16/07/2026", fontsize=10, color=red)
    doc.save(path)
    doc.close()
    return path


# --- split_list_breaks: non-list lines keep their line break ---------------


def test_split_list_breaks_keeps_non_list_line_breaks():
    d = docx.Document()
    p = d.add_paragraph()
    for k, text in enumerate(
        ["Dear Sir,", "Please find:", "1. First item that wraps", "onto the next line"]
    ):
        if k:
            p.add_run().add_break()
        p.add_run(text)
    fx.split_list_breaks(d)
    assert [q.text for q in d.paragraphs] == [
        "Dear Sir,\nPlease find:",
        "1. First item that wraps\nonto the next line",
    ]


# --- _coalesce_row: word space + per-piece formatting survive --------------


def test_coalesced_spans_keep_space_and_own_format(tmp_path):
    lay = fx.analyze_pdf(_label_value_pdf(str(tmp_path / "lv.pdf")))
    first = lay.rows[0][0].spans[0]
    assert first.text == "Nombre: Juan Perez"
    assert first.color == 0xCC0000
    d = docx.Document()
    d.add_paragraph("Nombre: Juan Perez")
    d.add_paragraph("Fecha: 16/07/2026")
    fx.merge_row_paragraphs(d, lay)
    assert len(d.paragraphs) == 1
    runs = [
        (r.text, bool(r.bold), str(r.font.color.rgb))
        for r in d.paragraphs[0].runs
        if r.text.strip()
    ]
    assert runs[:2] == [("Nombre:", True, "CC0000"), (" Juan Perez", False, "CC0000")]
    assert d.paragraphs[0].text == "Nombre: Juan Perez\tFecha: 16/07/2026"


# --- text-only rebuilds must not drop pictures -----------------------------


def test_merge_row_keeps_inline_picture(tmp_path):
    lay = fx.analyze_pdf(_label_value_pdf(str(tmp_path / "lv.pdf")))
    d = docx.Document()
    p = d.add_paragraph("Nombre: Juan Perez")
    p.add_run().add_picture(io.BytesIO(_png()), width=Pt(20))
    d.add_paragraph("Fecha: 16/07/2026")
    fx.merge_row_paragraphs(d, lay)
    assert _drawings(d.element.body) == 1
    assert len(d.paragraphs) == 2  # left as-is rather than rebuilt


def test_flatten_rebuild_keeps_inline_picture(tmp_path):
    pdf = pdf_fixtures.make_official_letter(str(tmp_path / "carta.pdf"))
    d = _pdf2docx(pdf, str(tmp_path / "carta.docx"))
    target = next(p for p in d.paragraphs if p.text == "Datos personales")
    target.add_run().add_picture(io.BytesIO(_png()), width=Pt(20))
    before = _drawings(d.element.body)  # fixture logo + ours
    fx.flatten_column_sections(d, fx.analyze_pdf(pdf))
    assert _drawings(d.element.body) == before
    cells = [c for t in d.tables for c in t.rows[0].cells]
    assert any("Datos personales" in c.text and _drawings(c._tc) for c in cells)


# --- flatten fallback: order kept, every cell ends with a paragraph --------


def test_flatten_fallback_cell_structure(tmp_path):
    pdf = pdf_fixtures.make_official_letter(str(tmp_path / "carta.pdf"))
    d = _pdf2docx(pdf, str(tmp_path / "carta.docx"))
    hosts = [
        el
        for el in d.element.body
        if el.tag == f"{NS}p" and el.find(f"{NS}pPr/{NS}sectPr") is not None
    ]
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "TBL-IN-RIGHT-COL"
    hosts[2].addprevious(t._tbl)  # table last in the right column
    # no PDF rows -> nothing matches -> fallback moves elements as-is
    fx.flatten_column_sections(d, fx.Layout(page_w=W, page_h=H, page_count=1, rows=[[]]))
    flat = next(tb for tb in d.tables if "Datos personales" in tb._tbl.xml)
    for cell in flat.rows[0].cells:
        kids = [k for k in cell._tc if k.tag != f"{NS}tcPr"]
        assert kids[-1].tag == f"{NS}p", "a cell must end with a paragraph"
        texts = [
            "".join(x.text or "" for x in k.iter(f"{NS}t")) for k in kids[:-1] if k.tag == f"{NS}p"
        ]
        assert "" not in texts, "no stray empty paragraph between items"
    left = [k for k in flat.rows[0].cells[0]._tc if k.tag == f"{NS}p"]
    assert "".join(x.text or "" for x in left[0].iter(f"{NS}t")) == "Datos personales"


# --- redaction: header text goes, a background image stays -----------------


def test_redaction_keeps_background_image(tmp_path):
    src, dst = str(tmp_path / "bg.pdf"), str(tmp_path / "bg_red.pdf")
    bg = _png((240, 240, 255), (50, 70))
    doc = fitz.open()
    for pg in range(3):
        page = doc.new_page(width=W, height=H)
        page.insert_image(page.rect, stream=bg)
        page.insert_text((40, 40), "ACME S.A. de C.V. - Membrete", fontsize=11)
        _body_lines(page, pg)
    doc.save(src)
    doc.close()
    lay = fx.analyze_pdf(src)
    assert lay.header is not None
    fx.redact_bands(src, dst, lay)
    with fitz.open(dst) as red:
        for page in red:
            assert len(page.get_images()) == 1
            assert "Membrete" not in page.get_text()
            assert "Parrafo" in page.get_text()


# --- rule-delimited header: varying folio stays body content ---------------


def test_rule_header_varying_folio_not_frozen(tmp_path):
    src = str(tmp_path / "folio.pdf")
    doc = fitz.open()
    for pg in range(4):
        page = doc.new_page(width=W, height=H)
        # separate blocks: an exact letterhead line and a per-page folio
        page.insert_text((40, 30), "NOTARIA 88", fontsize=11)
        page.insert_text((400, 52), f"Folio {1001 + pg}", fontsize=11)
        page.draw_line(fitz.Point(40, 60), fitz.Point(W - 40, 60), width=0.8)
        _body_lines(page, pg)
    doc.save(src)
    doc.close()
    lay = fx.analyze_pdf(src)
    assert lay.header is not None and lay.header.rule
    texts = [" ".join(s.text for ln in b.lines for s in ln) for b in lay.header.blocks]
    assert any("NOTARIA 88" in t for t in texts)
    assert not any("Folio" in t for t in texts)
    out = str(tmp_path / "folio.docx")
    documents.pdf_to_docx(src, out)
    body = "\n".join(p.text for p in docx.Document(out).paragraphs)
    for n in (1001, 1002, 1003, 1004):
        assert f"Folio {n}" in body


# --- titlePg: page 1 keeps the band it really has --------------------------


def test_first_page_footer_survives_title_page(tmp_path):
    src = str(tmp_path / "first.pdf")
    doc = fitz.open()
    for pg in range(4):
        page = doc.new_page(width=W, height=H)
        if pg:
            page.insert_text((40, 40), "ACME S.A. de C.V. - Membrete", fontsize=11)
        _body_lines(page, pg)
        page.insert_text((200, H - 30), "Documento confidencial ACME", fontsize=8)
    doc.save(src)
    doc.close()
    out = str(tmp_path / "first.docx")
    documents.pdf_to_docx(src, out)
    sec = docx.Document(out).sections[0]
    assert sec.different_first_page_header_footer
    assert "Membrete" in sec.header.tables[0].rows[0].cells[0].text
    assert not sec.first_page_header.tables  # page 1 had no letterhead
    first_footer = "".join(c.text for t in sec.first_page_footer.tables for c in t.rows[0].cells)
    assert "Documento confidencial ACME" in first_footer


# --- has_strangled_tables: a real narrow column is not "strangled" ---------


def test_narrow_lattice_column_keeps_real_widths(tmp_path):
    src = str(tmp_path / "narrow.pdf")
    cols = [40, 90, 400, 520]  # 50 / 310 / 120 pt
    rows = [
        ("Cant.", "Descripcion", "Importe"),
        ("1.00", "Servicio de asesoria legal", "1,500.00"),
        ("2.00", "Copias certificadas", "300.00"),
    ]
    doc = fitz.open()
    page = doc.new_page()
    y0 = 100
    for i in range(len(rows) + 1):
        page.draw_line((cols[0], y0 + 20 * i), (cols[-1], y0 + 20 * i))
    for x in cols:
        page.draw_line((x, y0), (x, y0 + 20 * len(rows)))
    for i, r in enumerate(rows):
        for k, t in enumerate(r):
            page.insert_text((cols[k] + 4, y0 + 20 * i + 14), t, fontsize=10)
    doc.save(src)
    doc.close()
    out = str(tmp_path / "narrow.docx")
    documents.pdf_to_docx(src, out)
    t = docx.Document(out).tables[0]
    widths = [int(g.get(f"{NS}w")) for g in t._tbl.tblGrid.findall(f"{NS}gridCol")]
    real = [(b - a) * 20 for a, b in pairwise(cols)]
    for got, want in zip(widths, real, strict=True):
        assert abs(got - want) <= 200, (widths, real)


# --- body_pages: column sections don't add pages --------------------------


def test_body_pages_ignore_column_section_breaks(tmp_path):
    one = pdf_fixtures.make_official_letter(str(tmp_path / "carta.pdf"))
    two = str(tmp_path / "carta2.pdf")
    with fitz.open() as d2, fitz.open(one) as src:
        d2.insert_pdf(src)
        d2.insert_pdf(src)
        d2.save(two)
    assert len(fx.body_pages(_pdf2docx(one, str(tmp_path / "1.docx")))) == 1
    assert len(fx.body_pages(_pdf2docx(two, str(tmp_path / "2.docx")))) == 2
