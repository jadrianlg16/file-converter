"""Layout fixup for PDF -> DOCX — companion to ``documents.pdf_to_docx``.

pdf2docx alone reproduces three classes of layout damage (all captured by
tests/pdf_fixtures.py):

* Repeated per-page letterhead/footer content is emitted as *body* content on
  every page (often as fake layout tables) instead of real Word
  headers/footers — so editing the body dislocates every "header" below it.
* Side-by-side text (label/value forms, left|center|right rows) becomes
  "stream tables" with invented column widths and cell indents that can leave
  a *negative* usable width (text renders one character per line).
* Separate short lines are merged into justified paragraphs even when the
  source was ragged-right.

This module analyses the source PDF with PyMuPDF and repairs the converted
document with python-docx:

``analyze_pdf``             geometry snapshot: visual rows, blocks, repeated
                            header/footer bands (text with digits wildcarded,
                            horizontal rule, repeated images/logos).
``redact_bands``            strip the repeated bands from a working copy of
                            the PDF before pdf2docx sees it.
``build_header_footer``     rebuild the bands as a real Word header/footer:
                            1x3 borderless table bucketed left/center/right
                            by x-position, PAGE/NUMPAGES fields for varying
                            page numbers, empty first-page header when the
                            band never appears on page 1.
``fix_stream_table_widths`` real column widths from PDF geometry + clamp
                            impossible cell indents.
``has_strangled_tables``    detect leftover invented-table pathology so the
                            caller can re-convert without stream tables.
``merge_row_paragraphs``    re-join same-visual-row paragraphs into one
                            paragraph with tab stops at the PDF x-positions.
``split_list_breaks``       split merged list-like lines (1. / a) / bullets)
                            back into separate paragraphs.
``fix_justified_ragged``    drop inferred justification when the source block
                            had a ragged right edge.

Everything is best-effort: helpers raise freely and the caller guards them,
falling back to the plain pdf2docx output.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field, replace

import fitz  # PyMuPDF
from docx.enum.table import WD_ROW_HEIGHT_RULE  # noqa: F401  (re-export convenience)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DIGITS = re.compile(r"\d+")
_EMU_PER_PT = 12700
_TWIPS_PER_PT = 20

# A row qualifies for tab-stop merging when neighbouring spans are separated
# by at least this gap (pt) — smaller gaps are just word spacing.
_ROW_GAP_PT = 12.0
# Bands: header candidates live entirely above 18% of the page height,
# footer candidates entirely below 82% (a detected horizontal rule overrides
# these — see _detect_bands). A band must also start close to the page edge
# and be contiguous (no gap larger than _BAND_GAP_PT between its blocks).
_HEADER_LIMIT = 0.18
_FOOTER_LIMIT = 0.82
_EDGE_START = 0.12
_BAND_GAP_PT = 24.0


def _norm_repeat(text: str) -> str:
    """Whitespace-collapsed text with digit runs wildcarded (page numbers,
    dates and folios compare equal across pages)."""
    return _DIGITS.sub("#", " ".join(text.split())).strip()


def _squash(text: str) -> str:
    return "".join(text.split())


# --------------------------------------------------------------------------
# PDF-side analysis
# --------------------------------------------------------------------------

@dataclass
class Span:
    text: str
    x0: float
    x1: float
    y0: float
    y1: float
    size: float
    font: str
    bold: bool
    italic: bool
    color: int = 0  # sRGB int as reported by PyMuPDF (0 = black)
    lead_space: bool = False   # raw text had leading/trailing whitespace
    trail_space: bool = False  # (stripped, but still a word boundary)
    parts: list = field(default_factory=list)  # coalesced pieces, own format


@dataclass
class Row:
    """Spans sharing one visual baseline, wide gaps between them."""
    spans: list
    squash: str


@dataclass
class BlockInfo:
    """One fitz text block: used for raggedness lookup."""
    squash: str
    line_x1: list
    width: float


@dataclass
class RepeatedBlock:
    lines: list                      # list[list[Span]] from first occurrence
    occurrences: dict = field(default_factory=dict)   # page index -> fitz.Rect
    span_markers: dict = field(default_factory=dict)  # (li, si) -> [(s, e, instr)]
    bbox: tuple = (0, 0, 0, 0)
    exact: bool = True               # repeats with identical text (digits too)
    has_page_token: bool = False     # contains a page-number-tracking digit


@dataclass
class Band:
    blocks: list = field(default_factory=list)
    rule: dict | None = None         # {"occurrences": {page: Rect}}
    images: list = field(default_factory=list)  # {occurrences, data, ext, bbox}
    on_first_page: bool = True


@dataclass
class Layout:
    page_w: float
    page_h: float
    page_count: int
    rows: list = field(default_factory=list)     # per page: list[Row]
    blocks: list = field(default_factory=list)   # per page: list[BlockInfo]
    vlines: list = field(default_factory=list)   # per page: [(x, y0, y1)]
    fillrects: list = field(default_factory=list)  # per page: [(Rect, rgb)]
    header: Band | None = None
    footer: Band | None = None


def _mk_span(sp: dict) -> Span | None:
    text = sp.get("text", "")
    if not text.strip():
        return None
    flags = sp.get("flags", 0)
    fname = sp.get("font", "") or ""
    x0, y0, x1, y1 = sp["bbox"]
    return Span(
        text=text.strip(), x0=x0, x1=x1, y0=y0, y1=y1,
        size=float(sp.get("size", 11.0)), font=fname,
        bold=bool(flags & 16) or "bold" in fname.lower(),
        italic=bool(flags & 2) or "italic" in fname.lower()
              or "oblique" in fname.lower(),
        color=int(sp.get("color", 0)),
        lead_space=text[:1].isspace(), trail_space=text[-1:].isspace(),
    )


def _coalesce_row(spans: list) -> list:
    """Merge spans on one baseline whose gap is just word spacing. The merged
    span keeps each piece in ``parts`` so runs retain their own format."""
    spans = sorted(spans, key=lambda s: s.x0)
    out: list[Span] = []
    for sp in spans:
        if out and sp.x0 - out[-1].x1 < max(6.0, out[-1].size * 0.6):
            prev = out[-1]
            gap = sp.x0 - prev.x1
            # PyMuPDF often keeps the word space inside a span's text (which
            # _mk_span strips), so the bbox gap alone can read as zero
            sep = (" " if gap > prev.size * 0.12 or prev.trail_space
                   or sp.lead_space else "")
            out[-1] = Span(
                text=prev.text + sep + sp.text, x0=prev.x0, x1=sp.x1,
                y0=min(prev.y0, sp.y0), y1=max(prev.y1, sp.y1),
                size=prev.size, font=prev.font, bold=prev.bold,
                italic=prev.italic, color=prev.color,
                trail_space=sp.trail_space,
                parts=(prev.parts or [prev])
                + [replace(sp, text=sep + sp.text)],
            )
        else:
            out.append(sp)
    return out


def _page_rows(page) -> list:
    """Group every span on the page into visual rows (baseline clusters)."""
    spans = []
    for blk in page.get_text("dict")["blocks"]:
        if blk.get("type") != 0:
            continue
        for ln in blk.get("lines", []):
            for sp in ln.get("spans", []):
                s = _mk_span(sp)
                if s:
                    spans.append(s)
    spans.sort(key=lambda s: (s.y0, s.x0))
    rows, current, cur_y = [], [], None
    for sp in spans:
        if cur_y is None or abs(sp.y0 - cur_y) <= 2.0:
            current.append(sp)
            cur_y = sp.y0 if cur_y is None else cur_y
        else:
            rows.append(current)
            current, cur_y = [sp], sp.y0
    if current:
        rows.append(current)

    out = []
    for group in rows:
        merged = _coalesce_row(group)
        out.append(Row(spans=merged,
                       squash=_squash("".join(s.text for s in merged))))
    return out


def _page_blocks(page) -> tuple:
    """(BlockInfo list for raggedness, raw block records for band detection)."""
    infos, records = [], []
    for blk in page.get_text("dict")["blocks"]:
        if blk.get("type") != 0:
            continue
        lines, x1s, texts = [], [], []
        for ln in blk.get("lines", []):
            row = [s for s in (_mk_span(sp) for sp in ln.get("spans", []))
                   if s]
            if not row:
                continue
            row = _coalesce_row(row)
            lines.append(row)
            x1s.append(max(s.x1 for s in row))
            texts.append(" ".join(s.text for s in row))
        if not lines:
            continue
        full = "\n".join(texts)
        x0, y0, x1, y1 = blk["bbox"]
        infos.append(BlockInfo(squash=_squash(full), line_x1=x1s,
                               width=x1 - x0))
        records.append({
            "norm": _norm_repeat(full), "texts": texts, "lines": lines,
            "rect": fitz.Rect(blk["bbox"]),
        })
    return infos, records


def _line_markers(first: str, others: list, page_nums: list,
                  page_count: int) -> list:
    """Positions in ``first`` whose digit runs should become PAGE/NUMPAGES
    fields. ``others``/``page_nums`` are the same line on other pages (raw
    text, 0-based page index; first entry corresponds to ``first``)."""
    tokens = list(_DIGITS.finditer(first))
    if not tokens:
        return []
    all_tokens = [list(_DIGITS.finditer(t)) for t in others]
    if any(len(t) != len(tokens) for t in all_tokens):
        return []
    markers, has_page = [], False
    for k, tok in enumerate(tokens):
        values = [int(t[k].group()) for t in all_tokens]
        if all(v == p + 1 for v, p in zip(values, page_nums)):
            if len(set(values)) > 1:      # genuinely varying page number
                markers.append((tok.start(), tok.end(), "PAGE"))
                has_page = True
    if has_page:
        for k, tok in enumerate(tokens):
            values = [int(t[k].group()) for t in all_tokens]
            if len(set(values)) == 1 and values[0] == page_count:
                markers.append((tok.start(), tok.end(), "NUMPAGES"))
    return sorted(markers)


def _span_markers(lines: list, line_markers: list) -> dict:
    """Convert line-level PAGE/NUMPAGES markers into span-local offsets."""
    out: dict = {}
    for li, spans in enumerate(lines):
        markers = line_markers[li] if li < len(line_markers) else []
        if not markers:
            continue
        offset = 0
        for si, sp in enumerate(spans):
            if si:
                offset += 1  # single joining space, matches _page_blocks texts
            start, end = offset, offset + len(sp.text)
            local = [(ms - start, me - start, instr)
                     for ms, me, instr in markers
                     if start <= ms and me <= end]
            if local:
                out[(li, si)] = local
            offset = end
    return out


def _find_rules(hlines_per_page: list, W: float, H: float,
                threshold: int) -> dict:
    """Thin horizontal lines repeated near the top/bottom of most pages —
    the natural separator between a letterhead and the body."""
    rules = {"header": None, "footer": None}
    occ: dict = {"header": {}, "footer": {}}
    for pno, hlines in enumerate(hlines_per_page):
        for r in hlines:
            if r.width < 0.4 * W:
                continue
            zone = ("header" if r.y1 <= 0.25 * H
                    else "footer" if r.y0 >= 0.78 * H else None)
            if zone and pno not in occ[zone]:
                occ[zone][pno] = fitz.Rect(r.x0 - 1, r.y0 - 1,
                                           r.x1 + 1, r.y1 + 1)
    for zone in ("header", "footer"):
        if len(occ[zone]) >= threshold:
            ys = sorted(r.y0 for r in occ[zone].values())
            median = ys[len(ys) // 2]
            if all(abs(r.y0 - median) <= 5 for r in occ[zone].values()):
                rules[zone] = {"occurrences": occ[zone], "y": median}
    return rules


def _detect_bands(doc, page_records: list, hlines_per_page: list,
                  W: float, H: float) -> tuple:
    """Find blocks repeated across pages that form the header/footer bands.

    A horizontal rule (when present on most pages) delimits the band exactly.
    Without one, candidates must sit inside a narrow edge band, repeat with
    *identical* text or carry a page-number-tracking digit, start close to
    the page edge, and be contiguous — this keeps repeated body content
    (templated forms, numbered headings) out of the header.
    """
    n = len(page_records)
    if n < 2:
        return None, None
    threshold = max(2, -(-n * 3 // 5))  # ceil(0.6 * n)
    rules = _find_rules(hlines_per_page, W, H, threshold)

    zone_limit = {
        "header": (rules["header"]["y"] + 2 if rules["header"]
                   else H * _HEADER_LIMIT),
        "footer": (rules["footer"]["y"] - 2 if rules["footer"]
                   else H * _FOOTER_LIMIT),
    }

    groups: dict = {}
    for pno, records in enumerate(page_records):
        for rec in records:
            r = rec["rect"]
            if r.y1 <= zone_limit["header"]:
                zone = "header"
            elif r.y0 >= zone_limit["footer"]:
                zone = "footer"
            else:
                continue
            key = (zone, rec["norm"])
            grp = groups.setdefault(key, [])
            # same block repeats: similar y, similar left or right edge
            if grp and not any(
                abs(g["rect"].y0 - r.y0) <= 5 and
                (abs(g["rect"].x0 - r.x0) <= 8 or abs(g["rect"].x1 - r.x1) <= 8)
                for g in grp
            ):
                continue
            if any(g["page"] == pno for g in grp):
                continue
            grp.append({"page": pno, "rect": r, "rec": rec})

    bands = {"header": Band(), "footer": Band()}
    for (zone, norm), grp in groups.items():
        if len(grp) < threshold or not norm:
            continue
        first = grp[0]["rec"]
        rb = RepeatedBlock(lines=first["lines"], bbox=tuple(grp[0]["rect"]))
        for g in grp:
            rb.occurrences[g["page"]] = g["rect"]
        rb.exact = len({" ".join("\n".join(g["rec"]["texts"]).split())
                        for g in grp}) == 1
        page_nums = [g["page"] for g in grp]
        line_markers = []
        for li, line_text in enumerate(first["texts"]):
            texts, ok = [], True
            for g in grp:
                t = g["rec"]["texts"]
                if li >= len(t):
                    ok = False
                    break
                texts.append(t[li])
            line_markers.append(
                _line_markers(line_text, texts, page_nums, len(page_records))
                if ok else [])
        rb.span_markers = _span_markers(rb.lines, line_markers)
        rb.has_page_token = any(
            instr == "PAGE" for ms in rb.span_markers.values()
            for _, _, instr in ms)
        bands[zone].blocks.append(rb)

    for zone in ("header", "footer"):
        band = bands[zone]
        if not band.blocks:
            continue
        # only trust exact repeats or page-number blocks, rule or not: any
        # other varying digits (folios, dates) would be frozen to the first
        # occurrence while the real per-page values get redacted
        band.blocks = [b for b in band.blocks
                       if b.exact or b.has_page_token]
        if not rules[zone]:
            # no rule to delimit the band: contiguous from the page edge
            if zone == "header":
                band.blocks.sort(key=lambda b: b.bbox[1])
                kept, reach = [], H * _EDGE_START
                for b in band.blocks:
                    if b.bbox[1] <= reach + _BAND_GAP_PT:
                        kept.append(b)
                        reach = max(reach, b.bbox[3])
                band.blocks = kept
            else:
                band.blocks.sort(key=lambda b: -b.bbox[3])
                kept, reach = [], H * (1 - _EDGE_START)
                for b in band.blocks:
                    if b.bbox[3] >= reach - _BAND_GAP_PT:
                        kept.append(b)
                        reach = min(reach, b.bbox[1])
                band.blocks = kept

    header = bands["header"] if bands["header"].blocks else None
    footer = bands["footer"] if bands["footer"].blocks else None
    if header:
        header.on_first_page = all(0 in b.occurrences for b in header.blocks)
        header.rule = rules["header"]
        header.images = _find_band_images(doc, W, H, threshold, "header",
                                          zone_limit["header"])
    if footer:
        footer.on_first_page = all(0 in b.occurrences for b in footer.blocks)
        footer.rule = rules["footer"]
        footer.images = _find_band_images(doc, W, H, threshold, "footer",
                                          zone_limit["footer"])
    return header, footer


def _find_band_images(doc, W: float, H: float, threshold: int,
                      zone: str, limit_y: float | None = None) -> list:
    """Images (logos) with the same xref+position in the band on most pages."""
    out = []
    if limit_y is None:
        limit_y = H * (_HEADER_LIMIT if zone == "header" else _FOOTER_LIMIT)
    try:
        seen: dict = {}
        for pno in range(doc.page_count):
            for img in doc[pno].get_images(full=True):
                xref = img[0]
                for r in doc[pno].get_image_rects(xref):
                    in_band = (r.y1 <= limit_y if zone == "header"
                               else r.y0 >= limit_y)
                    if not in_band:
                        continue
                    entry = seen.setdefault(xref, {"occurrences": {}, "rect": r})
                    if abs(entry["rect"].y0 - r.y0) <= 5:
                        entry["occurrences"][pno] = r
        for xref, entry in seen.items():
            if len(entry["occurrences"]) < threshold:
                continue
            pix = doc.extract_image(xref)
            if pix and pix.get("ext") in ("png", "jpeg", "jpg", "bmp", "gif"):
                out.append({"occurrences": entry["occurrences"],
                            "data": pix["image"], "ext": pix["ext"],
                            "bbox": tuple(entry["rect"])})
    except Exception:
        return []
    return out


def analyze_pdf(path: str) -> Layout:
    with fitz.open(path) as doc:
        first = doc[0].rect
        layout = Layout(page_w=first.width, page_h=first.height,
                        page_count=doc.page_count)
        uniform = all(abs(doc[p].rect.width - first.width) < 2 and
                      abs(doc[p].rect.height - first.height) < 2
                      for p in range(doc.page_count))
        page_records = []
        hlines_per_page = []
        for pno in range(doc.page_count):
            page = doc[pno]
            layout.rows.append(_page_rows(page))
            infos, records = _page_blocks(page)
            layout.blocks.append(infos)
            page_records.append(records)
            hlines, vlines, fills = [], [], []
            if pno < 30:  # drawings are only needed for rules/borders/banners
                try:
                    for d in page.get_drawings():
                        r = d.get("rect")
                        if r is None:
                            continue
                        if r.height <= 3 and r.width >= 20:
                            hlines.append(r)
                        elif r.width <= 3 and r.height >= 16:
                            vlines.append((r.x0, r.y0, r.y1))
                        elif (d.get("fill") is not None and 6 <= r.height <= 80
                              and r.width >= 30):
                            fills.append((r, d["fill"]))
                except Exception:
                    pass
            hlines_per_page.append(hlines)
            layout.vlines.append(vlines)
            layout.fillrects.append(fills)
        if uniform:
            layout.header, layout.footer = _detect_bands(
                doc, page_records, hlines_per_page, first.width, first.height)
        return layout


# --------------------------------------------------------------------------
# PDF-side redaction
# --------------------------------------------------------------------------

def redact_bands(src: str, dst: str, layout: Layout) -> None:
    """Remove the detected header/footer content from a copy of the PDF."""
    doc = fitz.open(src)
    try:
        graphics_flag = getattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_COVERED", None)
        image_flag = getattr(fitz, "PDF_REDACT_IMAGE_REMOVE", 2)
        image_none = getattr(fitz, "PDF_REDACT_IMAGE_NONE", 0)
        for pno in range(doc.page_count):
            page = doc[pno]
            rects, img_rects = [], []
            for band in (layout.header, layout.footer):
                if not band:
                    continue
                for blk in band.blocks:
                    r = blk.occurrences.get(pno)
                    if r:
                        rects.append(fitz.Rect(r.x0 - 1, r.y0 - 1,
                                               r.x1 + 1, r.y1 + 1))
                if band.rule:
                    r = band.rule["occurrences"].get(pno)
                    if r:
                        rects.append(r)
                for img in band.images:
                    r = img["occurrences"].get(pno)
                    if r:
                        img_rects.append(fitz.Rect(r.x0 - 1, r.y0 - 1,
                                                   r.x1 + 1, r.y1 + 1))
            # two passes: text/rule rects must leave every image alone (a
            # background or body image merely touching a band block would
            # otherwise be removed whole); only band logos remove images
            for batch, images in ((rects, image_none),
                                  (img_rects, image_flag)):
                if not batch:
                    continue
                for r in batch:
                    page.add_redact_annot(r)
                if graphics_flag is not None:
                    page.apply_redactions(images=images,
                                          graphics=graphics_flag)
                else:  # older PyMuPDF without the graphics parameter
                    page.apply_redactions(images=images)
        doc.save(dst, garbage=3, deflate=True)
    finally:
        doc.close()


# --------------------------------------------------------------------------
# DOCX-side helpers
# --------------------------------------------------------------------------

def _pt(length) -> float:
    return int(length) / _EMU_PER_PT if length is not None else 0.0


def _apply_span_format(run, span: Span) -> None:
    from docx.shared import RGBColor

    run.font.size = Pt(round(span.size * 2) / 2)
    run.bold = span.bold
    run.italic = span.italic
    name = _map_font(span.font)
    if name:
        run.font.name = name
    if span.color:
        run.font.color.rgb = RGBColor((span.color >> 16) & 0xFF,
                                      (span.color >> 8) & 0xFF,
                                      span.color & 0xFF)


def _add_span_runs(paragraph, span: Span, tail: str = "") -> None:
    """Runs for ``span``: one per coalesced piece, each in its own format."""
    pieces = span.parts or [span]
    for k, piece in enumerate(pieces):
        text = piece.text + (tail if k == len(pieces) - 1 else "")
        _apply_span_format(paragraph.add_run(text), piece)


_FONT_MAP = [
    ("times", "Times New Roman"), ("georgia", "Georgia"),
    ("garamond", "Garamond"), ("cambria", "Cambria"),
    ("calibri", "Calibri"), ("verdana", "Verdana"),
    ("tahoma", "Tahoma"), ("courier", "Courier New"),
    ("consolas", "Consolas"), ("arial", "Arial"),
    ("helvetica", "Arial"), ("helv", "Arial"),
]


def _map_font(pdf_font: str) -> str | None:
    low = (pdf_font or "").lower()
    for key, name in _FONT_MAP:
        if key in low:
            return name
    return None


def _add_field(paragraph, instr: str, shown: str, span: Span) -> None:
    """Append a fldSimple (PAGE / NUMPAGES) rendered like ``span``."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(round(span.size * 2))))
    rpr.append(sz)
    if span.bold:
        rpr.append(OxmlElement("w:b"))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = shown
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _emit_span(paragraph, span: Span, markers: list, trailing_space: bool) -> None:
    """Write one span into ``paragraph`` as formatted runs, substituting
    PAGE/NUMPAGES fields at the span-local marker offsets."""
    if not markers:
        _add_span_runs(paragraph, span, " " if trailing_space else "")
        return
    text = span.text
    cursor = 0
    for ms, me, instr in sorted(markers):
        if ms < cursor or me > len(text):
            continue
        if ms > cursor:
            r = paragraph.add_run(text[cursor:ms])
            _apply_span_format(r, span)
        _add_field(paragraph, instr, text[ms:me], span)
        cursor = me
    tail = text[cursor:] + (" " if trailing_space else "")
    if tail:
        r = paragraph.add_run(tail)
        _apply_span_format(r, span)


def _shrink_paragraph(p) -> None:
    """Make an (empty) paragraph take almost no vertical space."""
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(2)
    for run in p.runs:
        run.font.size = Pt(2)


def _set_table_borders(table, bottom_rule: bool) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    bottom = OxmlElement("w:bottom")
    if bottom_rule:
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "000000")
    else:
        bottom.set(qn("w:val"), "none")
    borders.append(bottom)
    tbl_pr.append(borders)


def _zero_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _bucket(block_bbox: tuple, W: float) -> str:
    cx = (block_bbox[0] + block_bbox[2]) / 2
    if cx < W * 0.4:
        return "left"
    if cx > W * 0.6:
        return "right"
    return "center"


_BUCKET_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _build_band_content(hf, band: Band, layout: Layout, section) -> float:
    """Build the band table inside a header/footer part. Returns the
    estimated rendered height (pt) so the caller can reserve page margin."""
    hf.is_linked_to_previous = False
    W = layout.page_w
    margin_l = _pt(section.left_margin)
    margin_r = _pt(section.right_margin)
    text_w = max(W - margin_l - margin_r, 100.0)

    # Bucket at SPAN level: a single fitz block may span the whole page width
    # (left text + centered title + right date in one block), so block-level
    # bucketing would dump everything into the center cell.
    buckets = {"left": [], "center": [], "right": []}
    for blk in band.blocks:
        for li, line in enumerate(blk.lines):
            for si, sp in enumerate(line):
                markers = blk.span_markers.get((li, si), [])
                name = _bucket((sp.x0, sp.y0, sp.x1, sp.y1), W)
                buckets[name].append((sp, markers))
    img_buckets = {"left": [], "center": [], "right": []}
    for img in band.images:
        img_buckets[_bucket(img["bbox"], W)].append(img)

    # group each bucket's spans into visual lines (same baseline) up front,
    # so column widths can come from real line widths — the docx text area
    # may be narrower than the original header extent after redaction
    line_groups = {}
    for name in ("left", "center", "right"):
        spans = sorted(buckets[name], key=lambda it: (it[0].y0, it[0].x0))
        groups: list = []
        current: list = []
        for item in spans:
            if current and abs(item[0].y0 - current[0][0].y0) > 2.5:
                groups.append(current)
                current = []
            current.append(item)
        if current:
            groups.append(current)
        line_groups[name] = groups

    def _need(name):
        widths = [max(sp.x1 for sp, _ in g) - min(sp.x0 for sp, _ in g)
                  for g in line_groups[name]]
        widths += [img["bbox"][2] - img["bbox"][0]
                   for img in img_buckets[name]]
        return max(widths) if widths else 0.0

    # the 10% + 8pt slack absorbs substituted-font width growth in Word/LO;
    # symmetric side columns (only needed when a centered cell exists) keep
    # the center cell truly page-centered
    need_l, need_c, need_r = _need("left"), _need("center"), _need("right")

    def _pad(w):
        return w * 1.10 + 8 if w else 24.0

    if need_c and need_l and need_r:
        left_w = right_w = min(max(_pad(need_l), _pad(need_r)),
                               text_w * 0.40)
    else:
        left_w = min(_pad(need_l), text_w - 48)
        right_w = min(_pad(need_r), text_w - left_w - 24)
    center_w = max(text_w - left_w - right_w, 24.0)

    table = hf.add_table(rows=1, cols=3, width=Emu(int(text_w * _EMU_PER_PT)))
    table.autofit = False
    widths = [left_w, center_w, right_w]
    for gc, w in zip(table._tbl.tblGrid.findall(f"{W_NS}gridCol"), widths):
        gc.set(qn("w:w"), str(int(w * _TWIPS_PER_PT)))
    _set_table_borders(table, bottom_rule=band.rule is not None)
    _zero_cell_margins(table)

    for name, cell, width in zip(("left", "center", "right"),
                                 table.rows[0].cells, widths):
        cell.width = Emu(int(width * _EMU_PER_PT))
        first = True
        for img in img_buckets[name]:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.alignment = _BUCKET_ALIGN[name]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            w_pt = img["bbox"][2] - img["bbox"][0]
            try:
                p.add_run().add_picture(io.BytesIO(img["data"]),
                                        width=Emu(int(w_pt * _EMU_PER_PT)))
            except Exception:
                continue
        for group in line_groups[name]:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.alignment = _BUCKET_ALIGN[name]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for gi, (sp, markers) in enumerate(group):
                _emit_span(p, sp, markers, gi < len(group) - 1)

    # header/footer parts start with one default empty paragraph — keep it
    # (Word wants a trailing w:p after a table) but make it invisible, and
    # move the table above it.
    anchor = hf.paragraphs[0]
    anchor._p.addprevious(table._tbl)
    _shrink_paragraph(anchor)

    # estimated height: tallest bucket (sum of its line heights) + images
    est = 0.0
    for name in ("left", "center", "right"):
        h = sum(max(sp.size for sp, _ in g) * 1.3 for g in line_groups[name])
        h += sum(img["bbox"][3] - img["bbox"][1] for img in img_buckets[name])
        est = max(est, h)
    return est + 5.0  # shrunk trailing paragraph + border


def _first_page_band(band: Band) -> Band | None:
    """The part of ``band`` present on page 1 (None when nothing is)."""
    blocks = [b for b in band.blocks if 0 in b.occurrences]
    images = [i for i in band.images if 0 in i["occurrences"]]
    rule = (band.rule if band.rule and 0 in band.rule["occurrences"]
            else None)
    if not blocks and not images and not rule:
        return None
    return Band(blocks=blocks, rule=rule, images=images)


def build_header_footer(doc, layout: Layout) -> None:
    """Rebuild the detected repeated bands as real Word header/footer on the
    first section (later sections inherit — pdf2docx never writes its own).
    Page margins grow to fit the band so it doesn't push body content."""
    section = doc.sections[0]
    header_h = footer_h = 0.0
    if layout.header:
        header_h = _build_band_content(section.header, layout.header,
                                       layout, section)
    if layout.footer:
        footer_h = _build_band_content(section.footer, layout.footer,
                                       layout, section)
    if any(b and not b.on_first_page for b in (layout.header, layout.footer)):
        # titlePg blanks BOTH first-page parts: rebuild whatever part of each
        # band page 1 does carry (already redacted from its body) there
        section.different_first_page_header_footer = True
        for band, hf in ((layout.header, section.first_page_header),
                         (layout.footer, section.first_page_footer)):
            first = _first_page_band(band) if band else None
            if first:
                _build_band_content(hf, first, layout, section)
    for sec in doc.sections:
        sec.header_distance = Pt(14)
        sec.footer_distance = Pt(14)
        if header_h:
            need = Pt(14 + header_h + 4)
            if sec.top_margin is None or sec.top_margin < need:
                sec.top_margin = need
        if footer_h:
            need = Pt(14 + footer_h + 4)
            if sec.bottom_margin is None or sec.bottom_margin < need:
                sec.bottom_margin = need


# --------------------------------------------------------------------------
# Body repairs
# --------------------------------------------------------------------------

def _starts_new_page(sect_pr) -> bool:
    """A sectPr starts a new page unless its type is continuous/nextColumn
    (pdf2docx emits those for multi-column zones within one source page)."""
    t = sect_pr.find(f"{W_NS}type")
    return t is None or t.get(qn("w:val")) in ("nextPage", "oddPage",
                                               "evenPage")


def _body_pages(doc) -> list:
    """Split top-level body elements into per-page groups (pdf2docx emits one
    page-breaking section per source page; continuous/column sections stay
    within their page). A sectPr's type says how its *own* section starts,
    so a section ends its page only when the NEXT sectPr breaks the page."""
    els = list(doc.element.body)
    sects = [el.find(f"{W_NS}pPr/{W_NS}sectPr") if el.tag == f"{W_NS}p"
             else el if el.tag == f"{W_NS}sectPr" else None for el in els]
    idx = [k for k, s in enumerate(sects) if s is not None]
    ends_page = {a for a, b in zip(idx, idx[1:])
                 if _starts_new_page(sects[b])}
    pages, current = [], []
    for k, el in enumerate(els):
        tag = el.tag
        if tag == f"{W_NS}p":
            current.append(el)
            if k in ends_page:
                pages.append(current)
                current = []
        elif tag == f"{W_NS}tbl":
            current.append(el)
        elif tag == f"{W_NS}sectPr":
            pages.append(current)
            current = []
    if current:
        pages.append(current)
    return pages


_OPAQUE_TAGS = (f"{W_NS}drawing", f"{W_NS}pict", f"{W_NS}hyperlink")


def _has_opaque(el) -> bool:
    """True when ``el`` holds content a text-only rebuild from PDF spans
    would drop (inline pictures, VML, hyperlinks)."""
    return next(el.iter(*_OPAQUE_TAGS), None) is not None


def _iter_container_paragraphs(container):
    """Every paragraph in a Document/_Cell, nested tables included."""
    for p in container.paragraphs:
        yield p
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_container_paragraphs(cell)


def _iter_elements_paragraphs(elements, doc):
    """Paragraph objects in the given body elements, table cells included."""
    for el in elements:
        if el.tag == f"{W_NS}p":
            yield Paragraph(el, doc)
        elif el.tag == f"{W_NS}tbl":
            for row in Table(el, doc).rows:
                for cell in row.cells:
                    yield from _iter_container_paragraphs(cell)


def flatten_column_sections(doc, layout: Layout) -> None:
    """Replace pdf2docx's multi-column sections with a real layout table.

    pdf2docx models side-by-side zones (e.g. personal data left, place/date
    right in official letters) as consecutive Word sections with
    ``w:cols num="2"`` joined by ``nextColumn`` breaks. Word and LibreOffice
    render these poorly — content reflows across the columns and the section
    boundaries often force page breaks. A borderless 1xN table with one
    column-flow per cell renders identically everywhere and is editable.

    Cell content is rebuilt from PDF geometry (one paragraph per source line,
    original formatting) when every source row can be matched; otherwise the
    existing elements are moved into the cells unchanged.
    """
    body = doc.element.body

    # parse the body into sections
    sections, current = [], []
    for el in list(body):
        tag = el.tag
        if tag == f"{W_NS}p":
            current.append(el)
            spr = el.find(f"{W_NS}pPr/{W_NS}sectPr")
            if spr is not None:
                sections.append({"els": current, "sect": spr, "host": el})
                current = []
        elif tag == f"{W_NS}tbl":
            current.append(el)
        elif tag == f"{W_NS}sectPr":
            sections.append({"els": current, "sect": el, "host": None})
            current = []
    if current:
        sections.append({"els": current, "sect": None, "host": None})

    def _cols(sec):
        if sec["sect"] is None:
            return 1
        c = sec["sect"].find(f"{W_NS}cols")
        if c is None:
            return 1
        return int(c.get(qn("w:num")) or 1)

    # page index per section. A page-breaking sectPr immediately followed by
    # a multi-column section is pdf2docx's columns machinery, not a real
    # source page break — the column zone lives on the same page.
    page_of, pno = [], 0
    for k, sec in enumerate(sections):
        page_of.append(pno)
        if sec["sect"] is None or not _starts_new_page(sec["sect"]):
            continue
        next_cols = _cols(sections[k + 1]) if k + 1 < len(sections) else 1
        if next_cols < 2:
            pno += 1

    # group consecutive multi-column sections
    i = 0
    groups = []
    while i < len(sections):
        if _cols(sections[i]) >= 2:
            j = i
            while j < len(sections) and _cols(sections[j]) >= 2:
                j += 1
            groups.append((i, j, page_of[i]))
            i = j
        else:
            i += 1

    for start, end, gpno in groups:
        group = sections[start:end]
        anchor = next((el for sec in group for el in sec["els"]), None)
        if anchor is None:
            continue
        # pdf2docx closes the zone *before* the columns with a typeless
        # (page-breaking) sectPr even though it's the same source page —
        # neutralize it so the flattened flow stays on one page
        if start > 0 and page_of[start - 1] == gpno:
            prev = sections[start - 1]["sect"]
            if prev is not None:
                t = prev.find(f"{W_NS}type")
                if t is None:
                    t = OxmlElement("w:type")
                    prev.insert(0, t)
                t.set(qn("w:val"), "continuous")
        n = len(group)
        # column widths from the first section's explicit w:col list
        widths_tw = []
        c = group[0]["sect"].find(f"{W_NS}cols")
        if c is not None:
            widths_tw = [int(col.get(qn("w:w")) or 0)
                         for col in c.findall(f"{W_NS}col")]
        sec_obj = doc.sections[0]
        text_tw = int(sec_obj.page_width - sec_obj.left_margin -
                      sec_obj.right_margin) // 635
        if len(widths_tw) != n or not all(widths_tw):
            widths_tw = [text_tw // n] * n

        table = doc.add_table(rows=1, cols=n)
        table.autofit = False
        _set_table_borders(table, bottom_rule=False)
        for gc, w in zip(table._tbl.tblGrid.findall(f"{W_NS}gridCol"),
                         widths_tw):
            gc.set(qn("w:w"), str(w))
        anchor.addprevious(table._tbl)

        rows = layout.rows[gpno] if gpno < len(layout.rows) else []
        for sec, cell, w in zip(group, table.rows[0].cells, widths_tw):
            cell.width = Emu(w * 635)
            texts = [p.text for p in _iter_elements_paragraphs(sec["els"], doc)
                     if p.text.strip()]
            sq = _squash("".join(texts))
            matched = [r for r in rows if r.squash and r.squash in sq]
            covered = sum(len(r.squash) for r in matched)
            if (matched and covered == len(sq)
                    and not any(_has_opaque(el) for el in sec["els"])):
                # rebuild from geometry: one paragraph per source line
                matched.sort(key=lambda r: r.spans[0].y0)
                cl_x0 = min(s.x0 for r in matched for s in r.spans)
                cl_x1 = max(s.x1 for r in matched for s in r.spans)
                cl_cx = (cl_x0 + cl_x1) / 2
                first, prev_y1 = True, None
                for r in matched:
                    p = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    pf = p.paragraph_format
                    pf.space_after = Pt(0)
                    if prev_y1 is not None:
                        pf.space_before = Pt(max(r.spans[0].y0 - prev_y1, 0))
                    prev_y1 = max(s.y1 for s in r.spans)
                    rw = max(s.x1 for s in r.spans) - min(s.x0 for s in r.spans)
                    rcx = (min(s.x0 for s in r.spans) +
                           max(s.x1 for s in r.spans)) / 2
                    if rw < (cl_x1 - cl_x0) * 0.85 and abs(rcx - cl_cx) <= 8:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for si, sp in enumerate(r.spans):
                        _add_span_runs(p, sp,
                                       " " if si < len(r.spans) - 1 else "")
                for el in sec["els"]:
                    body.remove(el)
            else:
                # fallback: move the existing elements into the cell, in
                # order, ahead of its empty placeholder paragraph — a cell
                # must end with a w:p, so the placeholder goes only when a
                # paragraph was moved last
                placeholder = cell.paragraphs[0]._p
                last = None
                for el in sec["els"]:
                    if el is sec["host"]:
                        ppr = el.find(f"{W_NS}pPr")
                        spr = (ppr.find(f"{W_NS}sectPr")
                               if ppr is not None else None)
                        if spr is not None:
                            ppr.remove(spr)
                        if not el.findall(f"{W_NS}r"):
                            body.remove(el)
                            continue
                    placeholder.addprevious(el)
                    last = el
                if last is not None and last.tag == f"{W_NS}p":
                    placeholder.getparent().remove(placeholder)
            # drop the section break itself (host handled above; body-level
            # sectPr for mid-document groups should not survive either)
            if sec["host"] is not None and sec["host"].getparent() is not None:
                ppr = sec["host"].find(f"{W_NS}pPr")
                spr = ppr.find(f"{W_NS}sectPr") if ppr is not None else None
                if spr is not None:
                    ppr.remove(spr)
                if not sec["host"].findall(f"{W_NS}r"):
                    body.remove(sec["host"])


def _is_light(color_val: int) -> bool:
    r, g, b = (color_val >> 16) & 0xFF, (color_val >> 8) & 0xFF, color_val & 0xFF
    return (r * 299 + g * 587 + b * 114) / 1000 > 200


def restore_banner_shading(doc, layout: Layout) -> None:
    """Light text drawn over a filled rectangle (e.g. white-on-red banner
    titles in official letters) loses the rectangle in pdf2docx and becomes
    invisible. Re-apply the fill as paragraph shading."""
    pages = _body_pages(doc)
    for pno, elements in enumerate(pages):
        fills = layout.fillrects[pno] if pno < len(layout.fillrects) else []
        if not fills:
            continue
        rows = layout.rows[pno] if pno < len(layout.rows) else []
        for para in _iter_elements_paragraphs(elements, doc):
            probe = _squash(para.text)
            if not probe:
                continue
            colors = [r.font.color.rgb for r in para.runs
                      if r.font.color and r.font.color.rgb is not None]
            if not colors or not all(
                    _is_light(int(str(c), 16)) for c in colors):
                continue
            row = next((r for r in rows if r.squash == probe
                        or r.squash.startswith(probe)), None)
            if row is None:
                continue
            cx = (row.spans[0].x0 + row.spans[-1].x1) / 2
            cy = (row.spans[0].y0 + row.spans[0].y1) / 2
            rect_fill = next((fill for rect, fill in fills
                              if rect.contains(fitz.Point(cx, cy))), None)
            if rect_fill is None:
                continue
            rgb = "%02X%02X%02X" % tuple(int(round(v * 255))
                                         for v in rect_fill[:3])
            ppr = para._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), rgb)
            ppr.append(shd)


def has_strangled_tables(doc) -> bool:
    """True when any table cell leaves text less than ~50pt of usable width
    (pdf2docx sometimes emits indents wider than the cell itself)."""
    for table in doc.tables:
        grid = [int(gc.get(qn("w:w")) or 0)
                for gc in table._tbl.tblGrid.findall(f"{W_NS}gridCol")]
        for row in table.rows:
            for idx, cell in enumerate(row.cells):
                col_w = grid[idx] if idx < len(grid) else 0
                if not col_w:
                    continue
                for p in cell.paragraphs:
                    if len(p.text.strip()) < 4:
                        continue
                    pf = p.paragraph_format
                    li = int(pf.left_indent or 0) // 635
                    ri = int(pf.right_indent or 0) // 635
                    usable = col_w - li - ri - 216  # minus default cell margins
                    # only indents strangle: a genuinely narrow column (qty,
                    # unit) is fine, and re-converting would undo its widths
                    if usable < 50 * _TWIPS_PER_PT and li + ri and (
                            usable < 0 or col_w - 216 >= 50 * _TWIPS_PER_PT):
                        return True
    return False


def _table_has_borders(tbl_el) -> bool:
    """True when any cell/table border is actually drawn — pdf2docx derived
    this table from ruled lines (lattice), so its column widths are real."""
    for tag in ("tcBorders", "tblBorders"):
        for borders in tbl_el.iter(f"{W_NS}{tag}"):
            for edge in borders:
                if edge.get(qn("w:val")) not in (None, "none", "nil"):
                    return True
    return False


def fix_stream_table_widths(doc, layout: Layout) -> None:
    """Give layout tables their real column boundaries from PDF geometry and
    clamp cell indents that don't fit the cell.

    Column boundaries come from, in order of preference: vertical ruled lines
    that span the table's rows (lattice tables), then midpoints between the
    matched spans' extents (borderless stream tables). Bordered tables with
    no matching vertical lines keep pdf2docx's widths."""
    pages = _body_pages(doc)
    for pno, elements in enumerate(pages):
        if pno >= len(layout.rows):
            break
        rows = [r for r in layout.rows[pno] if len(r.spans) >= 2]
        vlines = layout.vlines[pno] if pno < len(layout.vlines) else []
        for el in elements:
            if el.tag != f"{W_NS}tbl":
                continue
            table = Table(el, doc)
            ncols = len(table.columns)
            grid = table._tbl.tblGrid.findall(f"{W_NS}gridCol")
            col_texts = ["" for _ in range(ncols)]
            for row in table.rows:
                for idx, cell in enumerate(row.cells[:ncols]):
                    col_texts[idx] += _squash(cell.text)
            # text + logo table (letterheads): give the text all the room the
            # image doesn't need, so the title stops wrapping
            if len(table.rows) == 1 and ncols == 2:
                has_img = [any(True for _ in c._tc.iter(f"{W_NS}drawing"))
                           for c in table.rows[0].cells]
                if has_img.count(True) == 1:
                    img_idx = has_img.index(True)
                    if not col_texts[img_idx]:
                        total = sum(int(gc.get(qn("w:w")) or 0) for gc in grid)
                        img_w = max(total // 5, 1400)
                        widths_tw = [total - img_w, img_w]
                        if img_idx == 0:
                            widths_tw.reverse()
                        for gc, w in zip(grid, widths_tw):
                            gc.set(qn("w:w"), str(w))
                        for idx, cell in enumerate(table.rows[0].cells):
                            cell.width = Emu(widths_tw[idx] * 635)
                        _clamp_cell_indents(table)
                        continue
            # match PDF rows whose spans fall into this table's columns
            matched = [r for r in rows
                       if len(r.spans) == ncols and all(
                           _squash(sp.text) and _squash(sp.text) in col_texts[k]
                           for k, sp in enumerate(r.spans))]
            edges = None
            if matched:
                y_top = min(r.spans[0].y0 for r in matched)
                y_bot = max(r.spans[0].y1 for r in matched)
                borders = sorted(x for x, ly0, ly1 in vlines
                                 if ly0 <= y_top + 8 and ly1 >= y_bot - 8)
                distinct = []
                for x in borders:
                    if not distinct or x - distinct[-1] > 2:
                        distinct.append(x)
                if len(distinct) == ncols + 1:
                    edges = distinct
                elif not _table_has_borders(el):
                    bounds = []
                    for k in range(ncols):
                        x0 = min(r.spans[k].x0 for r in matched)
                        x1 = max(r.spans[k].x1 for r in matched)
                        bounds.append((x0, x1))
                    edges = [bounds[0][0]]
                    for k in range(1, ncols):
                        edges.append((bounds[k - 1][1] + bounds[k][0]) / 2)
                    edges.append(bounds[-1][1] + 6)
            if edges:
                widths_tw = [max(int((edges[k + 1] - edges[k]) *
                                     _TWIPS_PER_PT), 300)
                             for k in range(ncols)]
                for gc, w in zip(grid, widths_tw):
                    gc.set(qn("w:w"), str(w))
                for row in table.rows:
                    for idx, cell in enumerate(row.cells[:ncols]):
                        cell.width = Emu(widths_tw[idx] * 635)
            # clamp impossible indents regardless of match
            _clamp_cell_indents(table)


def _clamp_cell_indents(table) -> None:
    grid = [int(gc.get(qn("w:w")) or 0)
            for gc in table._tbl.tblGrid.findall(f"{W_NS}gridCol")]
    ncols = len(table.columns)
    for row in table.rows:
        for idx, cell in enumerate(row.cells[:ncols]):
            col_w = grid[idx] if idx < len(grid) else 0
            if not col_w:
                continue
            for p in cell.paragraphs:
                pf = p.paragraph_format
                li = int(pf.left_indent or 0) // 635
                ri = int(pf.right_indent or 0) // 635
                if li + ri and col_w - li - ri < col_w * 0.35:
                    pf.left_indent = Pt(0)
                    pf.right_indent = Pt(0)


def merge_row_paragraphs(doc, layout: Layout) -> None:
    """Re-join consecutive paragraphs that were one visual row in the PDF
    into a single paragraph with tab stops at the original x-positions."""
    pages = _body_pages(doc)
    for pno, elements in enumerate(pages):
        if pno >= len(layout.rows):
            break
        section = doc.sections[min(pno, len(doc.sections) - 1)]
        margin_l = _pt(section.left_margin)
        margin_r = _pt(section.right_margin)
        candidates = []
        for row in layout.rows[pno]:
            if len(row.spans) < 2:
                continue
            gaps = [row.spans[i + 1].x0 - row.spans[i].x1
                    for i in range(len(row.spans) - 1)]
            if max(gaps) >= _ROW_GAP_PT:
                candidates.append(row)
        if not candidates:
            continue

        paras = [Paragraph(el, doc) for el in elements
                 if el.tag == f"{W_NS}p" and
                 el.find(f"{W_NS}pPr/{W_NS}sectPr") is None]
        paras = [p for p in paras if p.text.strip()]
        used = set()
        i = 0
        while i < len(paras):
            acc = _squash(paras[i].text)
            hit = None
            for ridx, row in enumerate(candidates):
                if ridx in used or not row.squash.startswith(acc):
                    continue
                j, a = i + 1, acc
                while (len(a) < len(row.squash) and j < len(paras)
                       and j - i < 6):
                    nxt = a + _squash(paras[j].text)
                    if not row.squash.startswith(nxt):
                        break
                    a = nxt
                    j += 1
                if a == row.squash and j - i >= 2:
                    hit = (ridx, row, j)
                    break
            # the merged paragraph is rebuilt from PDF text only
            if not hit or any(_has_opaque(p._p) for p in paras[i:hit[2]]):
                i += 1
                continue
            ridx, row, j = hit
            used.add(ridx)
            first = paras[i]
            new_p = OxmlElement("w:p")
            first._p.addprevious(new_p)
            merged = Paragraph(new_p, doc)
            pf = merged.paragraph_format
            pf.space_before = first.paragraph_format.space_before
            pf.space_after = Pt(0)
            for k, sp in enumerate(row.spans):
                if k == 0:
                    if sp.x0 - margin_l > 3:
                        pf.tab_stops.add_tab_stop(Pt(sp.x0 - margin_l),
                                                  WD_TAB_ALIGNMENT.LEFT)
                        merged.add_run("\t")
                else:
                    right_edge = layout.page_w - margin_r
                    if abs(sp.x1 - right_edge) <= 8:
                        pf.tab_stops.add_tab_stop(Pt(right_edge - margin_l),
                                                  WD_TAB_ALIGNMENT.RIGHT)
                    else:
                        pf.tab_stops.add_tab_stop(Pt(sp.x0 - margin_l),
                                                  WD_TAB_ALIGNMENT.LEFT)
                    merged.add_run("\t")
                _add_span_runs(merged, sp)
            for p in paras[i:j]:
                p._p.getparent().remove(p._p)
            paras[i:j] = [merged]
            i += 1


_LIST_START = re.compile(
    r"^\s*(\d{1,3}[.)°:]\s|[IVXLC]{1,6}[.)]\s|[a-z][.)]\s|[•▪◦‣∙*-]\s)",
    re.IGNORECASE)


def split_list_breaks(doc) -> None:
    """Split paragraphs at hard line breaks that start a list-like item."""
    for el in list(doc.element.body):
        if el.tag != f"{W_NS}p":
            continue
        para = Paragraph(el, doc)
        if "\n" not in para.text:
            continue
        # normalise: make every w:br sit in its own run
        for r in list(el.findall(f"{W_NS}r")):
            children = list(r)
            brs = [c for c in children if c.tag == f"{W_NS}br"]
            if not brs or len(children) <= 1:
                continue
            rpr = r.find(f"{W_NS}rPr")
            groups, cur = [], []
            for c in children:
                if c.tag == f"{W_NS}rPr":
                    continue
                if c.tag == f"{W_NS}br":
                    groups.append(cur)
                    groups.append("BR")
                    cur = []
                else:
                    cur.append(c)
            groups.append(cur)
            prev = r
            for g in groups:
                nr = OxmlElement("w:r")
                if rpr is not None:
                    nr.append(copy.deepcopy(rpr))
                if g == "BR":
                    nr.append(OxmlElement("w:br"))
                else:
                    for c in g:
                        nr.append(c)
                prev.addnext(nr)
                prev = nr
            r.getparent().remove(r)

        # partition runs at break-runs
        runs = [c for c in el if c.tag == f"{W_NS}r"]
        segments, breaks, cur = [], [], []
        for r in runs:
            if r.find(f"{W_NS}br") is not None and len(
                    [c for c in r if c.tag != f"{W_NS}rPr"]) == 1:
                segments.append(cur)
                breaks.append(r)  # breaks[k] precedes segments[k + 1]
                cur = []
            else:
                cur.append(r)
        segments.append(cur)
        if len(segments) < 2:
            continue

        def _seg_text(seg):
            return "".join(t.text or "" for r in seg
                           for t in r.findall(f"{W_NS}t"))

        splittable = [s for s in segments[1:] if _LIST_START.match(_seg_text(s))]
        if not splittable:
            continue

        ppr = el.find(f"{W_NS}pPr")
        created = [el]
        current_target = el
        dropped = []
        for seg, brk in zip(segments[1:], breaks):
            if seg and _LIST_START.match(_seg_text(seg)):
                dropped.append(brk)
                np = OxmlElement("w:p")
                if ppr is not None:
                    nppr = copy.deepcopy(ppr)
                    s = nppr.find(f"{W_NS}sectPr")
                    if s is not None:
                        nppr.remove(s)
                    np.append(nppr)
                created[-1].addnext(np)
                created.append(np)
                current_target = np
                for r in seg:
                    np.append(r)
            else:
                # a non-list line keeps its own line break (dropping it would
                # glue its first word onto the previous line)
                for r in [brk] + seg:
                    current_target.append(r)
        # remove the now-orphaned breaks that preceded the split-off items
        for r in dropped:
            r.getparent().remove(r)
        # a section break must stay on the LAST paragraph of the page
        sect = ppr.find(f"{W_NS}sectPr") if ppr is not None else None
        if sect is not None and created[-1] is not el:
            last = created[-1]
            last_ppr = last.find(f"{W_NS}pPr")
            if last_ppr is None:
                last_ppr = OxmlElement("w:pPr")
                last.insert(0, last_ppr)
            ppr.remove(sect)
            last_ppr.append(sect)
        # every split line fits one source line: left-align at full width
        for pel in created:
            p = Paragraph(pel, doc)
            p.paragraph_format.right_indent = Pt(0)
            if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def fix_justified_ragged(doc, layout: Layout) -> None:
    """pdf2docx infers JUSTIFY for merged blocks; when the source block's
    right edge was ragged, restore LEFT alignment and full usable width."""
    pages = _body_pages(doc)
    for pno, elements in enumerate(pages):
        if pno >= len(layout.blocks):
            break
        blocks = layout.blocks[pno]
        for el in elements:
            if el.tag != f"{W_NS}p":
                continue
            para = Paragraph(el, doc)
            pf = para.paragraph_format
            if pf.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                continue
            probe = _squash(para.text)[:32]
            if not probe:
                continue
            blk = next((b for b in blocks if b.squash.startswith(probe)
                        or probe.startswith(b.squash[:32])), None)
            if blk is None or len(blk.line_x1) < 3:
                continue
            interior = blk.line_x1[:-1]
            ragged = (max(interior) - min(interior)) > max(
                6.0, 0.03 * blk.width)
            if ragged:
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.right_indent = Pt(0)
