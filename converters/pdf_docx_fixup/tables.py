"""Table repairs: real column widths from PDF geometry, impossible indents."""

from __future__ import annotations

from typing import Any

from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from docx.table import Table

from .docx_xml import EMU_PER_TWIP, TWIPS_PER_PT, W_NS, body_pages
from .model import Layout, Row, squash

_DEFAULT_CELL_MARGINS_TW = 216  # Word's left + right cell padding


def _grid_widths(table: Table) -> list[int]:
    return [int(gc.get(qn("w:w")) or 0) for gc in table._tbl.tblGrid.findall(f"{W_NS}gridCol")]


def has_strangled_tables(doc: Any) -> bool:
    """True when any table cell leaves text less than ~50pt of usable width
    (pdf2docx sometimes emits indents wider than the cell itself)."""
    min_tw = 50 * TWIPS_PER_PT
    for table in doc.tables:
        grid = _grid_widths(table)
        for row in table.rows:
            for idx, cell in enumerate(row.cells):
                col_w = grid[idx] if idx < len(grid) else 0
                if not col_w:
                    continue
                for p in cell.paragraphs:
                    if len(p.text.strip()) < 4:
                        continue
                    pf = p.paragraph_format
                    li = int(pf.left_indent or 0) // EMU_PER_TWIP
                    ri = int(pf.right_indent or 0) // EMU_PER_TWIP
                    usable = col_w - li - ri - _DEFAULT_CELL_MARGINS_TW
                    # only indents strangle: a genuinely narrow column (qty,
                    # unit) is fine, and re-converting would undo its widths
                    if (
                        usable < min_tw
                        and li + ri
                        and (usable < 0 or col_w - _DEFAULT_CELL_MARGINS_TW >= min_tw)
                    ):
                        return True
    return False


def _table_has_borders(tbl_el: Any) -> bool:
    """True when any cell/table border is actually drawn — pdf2docx derived
    this table from ruled lines (lattice), so its column widths are real."""
    for tag in ("tcBorders", "tblBorders"):
        for borders in tbl_el.iter(f"{W_NS}{tag}"):
            for edge in borders:
                if edge.get(qn("w:val")) not in (None, "none", "nil"):
                    return True
    return False


def _clamp_cell_indents(table: Table) -> None:
    """Drop cell paragraph indents that would leave under 35% of the column."""
    grid = _grid_widths(table)
    ncols = len(table.columns)
    for row in table.rows:
        for idx, cell in enumerate(row.cells[:ncols]):
            col_w = grid[idx] if idx < len(grid) else 0
            if not col_w:
                continue
            for p in cell.paragraphs:
                pf = p.paragraph_format
                li = int(pf.left_indent or 0) // EMU_PER_TWIP
                ri = int(pf.right_indent or 0) // EMU_PER_TWIP
                if li + ri and col_w - li - ri < col_w * 0.35:
                    pf.left_indent = Pt(0)
                    pf.right_indent = Pt(0)


def _widen_text_beside_logo(table: Table, grid: list[Any], col_texts: list[str]) -> bool:
    """For a one-row text + logo table (a letterhead), give the text all the
    room the image doesn't need, so the title stops wrapping. Returns True
    when the table was one."""
    if len(table.rows) != 1 or len(col_texts) != 2:
        return False
    has_img = [any(True for _ in c._tc.iter(f"{W_NS}drawing")) for c in table.rows[0].cells]
    if has_img.count(True) != 1:
        return False
    img_idx = has_img.index(True)
    if col_texts[img_idx]:
        return False
    total = sum(int(gc.get(qn("w:w")) or 0) for gc in grid)
    img_w = max(total // 5, 1400)
    widths_tw = [total - img_w, img_w]
    if img_idx == 0:
        widths_tw.reverse()
    for gc, w in zip(grid, widths_tw, strict=False):
        gc.set(qn("w:w"), str(w))
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = Emu(widths_tw[idx] * EMU_PER_TWIP)
    return True


def _column_edges(
    matched: list[Row], vlines: list[tuple[float, float, float]], ncols: int, tbl_el: Any
) -> list[float] | None:
    """x-positions of the table's ncols + 1 column boundaries, or None.

    Prefers vertical ruled lines spanning the matched rows (lattice tables),
    then midpoints between the spans' extents (borderless stream tables).
    A bordered table without matching vertical lines keeps its widths."""
    y_top = min(r.spans[0].y0 for r in matched)
    y_bot = max(r.spans[0].y1 for r in matched)
    borders = sorted(x for x, ly0, ly1 in vlines if ly0 <= y_top + 8 and ly1 >= y_bot - 8)
    distinct: list[float] = []
    for x in borders:
        if not distinct or x - distinct[-1] > 2:
            distinct.append(x)
    if len(distinct) == ncols + 1:
        return distinct
    if _table_has_borders(tbl_el):
        return None
    bounds = [
        (min(r.spans[k].x0 for r in matched), max(r.spans[k].x1 for r in matched))
        for k in range(ncols)
    ]
    edges = [bounds[0][0]]
    for k in range(1, ncols):
        edges.append((bounds[k - 1][1] + bounds[k][0]) / 2)
    edges.append(bounds[-1][1] + 6)
    return edges


def _fix_table(table: Table, rows: list[Row], vlines: list[tuple[float, float, float]]) -> None:
    ncols = len(table.columns)
    grid = table._tbl.tblGrid.findall(f"{W_NS}gridCol")
    col_texts = ["" for _ in range(ncols)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells[:ncols]):
            col_texts[idx] += squash(cell.text)
    if _widen_text_beside_logo(table, grid, col_texts):
        _clamp_cell_indents(table)
        return
    # PDF rows whose spans fall into this table's columns
    matched = [
        r
        for r in rows
        if len(r.spans) == ncols
        and all(squash(sp.text) and squash(sp.text) in col_texts[k] for k, sp in enumerate(r.spans))
    ]
    edges = _column_edges(matched, vlines, ncols, table._tbl) if matched else None
    if edges:
        widths_tw = [max(int((edges[k + 1] - edges[k]) * TWIPS_PER_PT), 300) for k in range(ncols)]
        for gc, w in zip(grid, widths_tw, strict=False):
            gc.set(qn("w:w"), str(w))
        for row in table.rows:
            for idx, cell in enumerate(row.cells[:ncols]):
                cell.width = Emu(widths_tw[idx] * EMU_PER_TWIP)
    # clamp impossible indents regardless of match
    _clamp_cell_indents(table)


def fix_stream_table_widths(doc: Any, layout: Layout) -> None:
    """Give layout tables their real column boundaries from PDF geometry and
    clamp cell indents that don't fit the cell.

    Column boundaries come from, in order of preference: vertical ruled lines
    that span the table's rows (lattice tables), then midpoints between the
    matched spans' extents (borderless stream tables). Bordered tables with
    no matching vertical lines keep pdf2docx's widths."""
    for pno, elements in enumerate(body_pages(doc)):
        if pno >= len(layout.rows):
            break
        rows = [r for r in layout.rows[pno] if len(r.spans) >= 2]
        vlines = layout.vlines[pno] if pno < len(layout.vlines) else []
        for el in elements:
            if el.tag == f"{W_NS}tbl":
                _fix_table(Table(el, doc), rows, vlines)
