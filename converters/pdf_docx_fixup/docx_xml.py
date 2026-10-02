"""Shared python-docx / WordprocessingML helpers for the DOCX-side repairs."""
from __future__ import annotations

from collections.abc import Iterable, Iterator
from itertools import pairwise
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Length, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .model import Span

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
EMU_PER_PT = 12700
EMU_PER_TWIP = 635
TWIPS_PER_PT = 20

_FONT_MAP = [
    ("times", "Times New Roman"), ("georgia", "Georgia"),
    ("garamond", "Garamond"), ("cambria", "Cambria"),
    ("calibri", "Calibri"), ("verdana", "Verdana"),
    ("tahoma", "Tahoma"), ("courier", "Courier New"),
    ("consolas", "Consolas"), ("arial", "Arial"),
    ("helvetica", "Arial"), ("helv", "Arial"),
]

# Content a text-only rebuild from PDF spans would drop.
_OPAQUE_TAGS = (f"{W_NS}drawing", f"{W_NS}pict", f"{W_NS}hyperlink")


def pt(length: Length | None) -> float:
    """A python-docx length in points (0 for an unset length)."""
    return int(length) / EMU_PER_PT if length is not None else 0.0


def _map_font(pdf_font: str) -> str | None:
    low = (pdf_font or "").lower()
    for key, name in _FONT_MAP:
        if key in low:
            return name
    return None


def apply_span_format(run: Run, span: Span) -> None:
    """Give ``run`` the size, weight, slant, font and colour of ``span``."""
    run.font.size = Pt(round(span.size * 2) / 2)
    run.bold = span.bold
    run.italic = span.italic
    name = _map_font(span.font)
    if name:
        run.font.name = name
    if span.color:
        run.font.color.rgb = RGBColor((span.color >> 16) & 0xFF,
                                      (span.color >> 8) & 0xFF,
                                      span.color & 0xFF)


def add_span_runs(paragraph: Paragraph, span: Span, tail: str = "") -> None:
    """Runs for ``span``: one per coalesced piece, each in its own format."""
    pieces = span.parts or [span]
    for k, piece in enumerate(pieces):
        text = piece.text + (tail if k == len(pieces) - 1 else "")
        apply_span_format(paragraph.add_run(text), piece)


def set_table_borders(table: Table, bottom_rule: bool) -> None:
    """Hide every border of ``table`` except, optionally, a bottom rule."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    bottom = OxmlElement("w:bottom")
    if bottom_rule:
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "000000")
    else:
        bottom.set(qn("w:val"), "none")
    borders.append(bottom)
    tbl_pr.append(borders)


def starts_new_page(sect_pr: Any) -> bool:
    """A sectPr starts a new page unless its type is continuous/nextColumn
    (pdf2docx emits those for multi-column zones within one source page)."""
    t = sect_pr.find(f"{W_NS}type")
    return t is None or t.get(qn("w:val")) in ("nextPage", "oddPage",
                                               "evenPage")


def body_pages(doc: Any) -> list[list[Any]]:
    """Split top-level body elements into per-page groups (pdf2docx emits one
    page-breaking section per source page; continuous/column sections stay
    within their page). A sectPr's type says how its *own* section starts,
    so a section ends its page only when the NEXT sectPr breaks the page."""
    els = list(doc.element.body)
    sects = [el.find(f"{W_NS}pPr/{W_NS}sectPr") if el.tag == f"{W_NS}p"
             else el if el.tag == f"{W_NS}sectPr" else None for el in els]
    idx = [k for k, s in enumerate(sects) if s is not None]
    ends_page = {a for a, b in pairwise(idx) if starts_new_page(sects[b])}
    pages, current = [], []
    for k, el in enumerate(els):
        tag = el.tag
        if tag == f"{W_NS}p":
            current.append(el)
            if k in ends_page:
                pages.append(current)
                current = []
        elif tag == f"{W_NS}tbl":
            current.append(el)
        elif tag == f"{W_NS}sectPr":
            pages.append(current)
            current = []
    if current:
        pages.append(current)
    return pages


def has_opaque(el: Any) -> bool:
    """True when ``el`` holds content a text-only rebuild from PDF spans
    would drop (inline pictures, VML, hyperlinks)."""
    return next(el.iter(*_OPAQUE_TAGS), None) is not None


def _iter_container_paragraphs(container: Any) -> Iterator[Paragraph]:
    """Every paragraph in a Document/_Cell, nested tables included."""
    yield from container.paragraphs
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_container_paragraphs(cell)


def iter_elements_paragraphs(elements: Iterable[Any], doc: Any) -> Iterator[Paragraph]:
    """Paragraph objects in the given body elements, table cells included."""
    for el in elements:
        if el.tag == f"{W_NS}p":
            yield Paragraph(el, doc)
        elif el.tag == f"{W_NS}tbl":
            for row in Table(el, doc).rows:
                for cell in row.cells:
                    yield from _iter_container_paragraphs(cell)


def cell_paragraph(cell: _Cell, index: int) -> Paragraph:
    """The cell's ``index``-th paragraph to fill: its built-in empty one
    first, then new ones appended after it."""
    return cell.paragraphs[0] if index == 0 else cell.add_paragraph()
